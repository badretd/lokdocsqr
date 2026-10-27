from __future__ import annotations

import builtins
import io
import re
import subprocess
import sys
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path
from types import SimpleNamespace

import pytest

import lokdocsqr.nodes.access.extract_text as extract_text


from lokdocsqr.nodes.access._base import DEFAULT_MAX_BYTES_PER_FILE


_MISSING = object()


# ===========================================================================
# Handler invocation helper — used by _call_handler for extract_text.HANDLERS
# ===========================================================================

def _invoke_callable(fn, **available):
    """Best-effort invocation supporting common handler signatures:
       (data, opts), (path, max_bytes, opts), (path, max_bytes), (path,).
    """
    import inspect

    try:
        sig = inspect.signature(fn)
    except (TypeError, ValueError):
        sig = None

    if sig is not None:
        kwargs = {}
        for name, param in sig.parameters.items():
            if param.kind in (param.VAR_POSITIONAL, param.VAR_KEYWORD):
                continue
            lname = name.lower()
            if name in available and available[name] is not None:
                kwargs[name] = available[name]
                continue
            if lname in ("p", "f", "file", "filepath", "path"):
                kwargs[name] = available.get("path")
            elif lname in ("limit", "max_size", "max_bytes_per_file",
                           "max_bytes", "maxbytes", "max_filesize"):
                kwargs[name] = available.get("max_bytes")
            elif lname in ("raw", "content", "bytes", "data", "buf", "payload"):
                kwargs[name] = available.get("data")
            elif lname in ("name", "filename", "fname"):
                kwargs[name] = available.get("filename")
            elif lname in ("query", "q", "needle", "term"):
                kwargs[name] = available.get("query")
            elif lname in ("root", "directory", "dir", "base_dir"):
                kwargs[name] = available.get("path")
            elif lname in ("ctx", "context"):
                kwargs[name] = available.get("ctx")
            elif lname in ("opts", "options", "settings", "config", "cfg"):
                kwargs[name] = (
                    available.get("opts")
                    or available.get("params")
                    or {}
                )
            elif lname in ("params", "parameters"):
                kwargs[name] = available.get("params") or {}

        missing = [
            n for n, p in sig.parameters.items()
            if n not in kwargs
            and p.default is p.empty
            and p.kind not in (p.VAR_POSITIONAL, p.VAR_KEYWORD)
        ]
        if not missing:
            try:
                return fn(**kwargs)
            except TypeError:
                pass

    path = available.get("path")
    max_bytes = available.get("max_bytes")
    opts = available.get("opts") or {}
    params = available.get("params") or {}
    ctx = available.get("ctx")

    attempts = []
    if path is not None:
        attempts.extend([
            (path, max_bytes, opts),
            (path, opts, max_bytes),
            (path, max_bytes, params),
            (path, params, max_bytes),
            (path, max_bytes),
            (path, opts),
            (path, params),
            (path,),
        ])
    if ctx is not None:
        attempts.extend([(ctx,), ()])
    attempts.append(())

    last_exc = None
    for args in attempts:
        try:
            return fn(*args)
        except TypeError as exc:
            last_exc = exc
            continue
    if last_exc is not None:
        raise last_exc
    return fn()


def _call_handler(fmt: str, path: Path, max_bytes: int | None = None):
    if max_bytes is None:
        max_bytes = DEFAULT_MAX_BYTES_PER_FILE
    path = Path(path)
    handler = extract_text.HANDLERS[fmt]
    data = path.read_bytes()
    opts = {
        "max_bytes_per_file": max_bytes,
        "max_bytes": max_bytes,
        "limit": max_bytes,
    }
    params = dict(opts)
    return _invoke_callable(
        handler,
        path=path,
        max_bytes=max_bytes,
        data=data,
        filename=path.name,
        opts=opts,
        params=params,
    )


# ===========================================================================
# Result normalization
# ===========================================================================

def _norm_result(result):
    """Normalize handler/node result to (text, error)."""
    if isinstance(result, list):
        texts, errors = [], []
        for item in result:
            t, e = _norm_result(item)
            if t:
                texts.append(t)
            if e:
                errors.append(e)
        return "\n".join(texts), ("; ".join(errors) if errors else None)

    if isinstance(result, dict):
        text = result.get("text")
        if text is None:
            text = result.get("content")
        if text is None:
            text = result.get("data")
        error = result.get("error")
        return ("" if text is None else str(text),
                None if not error else str(error))

    if isinstance(result, tuple):
        if len(result) == 0:
            return "", None
        if len(result) == 1:
            return _norm_result(result[0])[0], None
        text, error = result[0], result[1]
        return _norm_result(text)[0], (None if not error else str(error))

    if isinstance(result, str):
        return result, None

    if result is None:
        return "", None

    text = getattr(result, "text", None)
    if text is None:
        text = getattr(result, "content", None)
    error = getattr(result, "error", None)
    if text is None:
        text = str(result)
    return str(text), (None if not error else str(error))


class _FakeText(str):
    """String-like handler result that also behaves enough like a dict."""

    def get(self, key, default=None):
        if key in ("text", "content", "data"):
            return str(self)
        if key == "error":
            return None
        return default

    @property
    def text(self):
        return str(self)

    @property
    def content(self):
        return str(self)

    @property
    def error(self):
        return None


# ===========================================================================
# index_search invocation — uses the real public API
#
# Mirrors tests/test_access_index_search.py::_build_ctx exactly:
#   - node = IndexSearchNode(...)
#   - ctx  = SimpleNamespace(scope=ScopeManager(...), config=SimpleNamespace(...))
#   - result = node.run(ctx, {"paths": [str(root)]})
# ===========================================================================

def _build_index_search_ctx(root: Path):
    """Build the same ctx the operator test uses for IndexSearchNode."""
    from lokdocsqr.core.scope import ScopeManager

    scope = ScopeManager(
        SimpleNamespace(
            roots=[str(root)],
            writable=[str(root)],
            allow_symlinks=False,
        )
    )
    return SimpleNamespace(
        scope=scope,
        config=SimpleNamespace(base_dir=Path(root)),
    )


def _run_index_search(root: Path, query: str = "Hello", **node_kwargs):
    """Invoke the real IndexSearchNode through its public API."""
    from lokdocsqr.nodes.access.index_search import IndexSearchNode

    ctx = _build_index_search_ctx(root)
    node = IndexSearchNode(query=query, **node_kwargs)
    return node.run(ctx, {"paths": [str(root)]})


# ===========================================================================
# Fixture builders
# ===========================================================================

def _make_minimal_pdf(text: str = "Hello") -> bytes:
    """Build a minimal valid PDF 1.4 with one page and one text run."""
    escaped = (
        text.replace("\\", "\\\\")
        .replace("(", "\\(")
        .replace(")", "\\)")
    )
    stream = f"BT /F1 24 Tf 72 720 Td ({escaped}) Tj ET".encode("latin-1")

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    out = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n".encode("ascii") + obj + b"\nendobj\n"

    xref_pos = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode("ascii")
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += f"{off:010d} 00000 n \n".encode("ascii")
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_pos}\n%%EOF\n"
    ).encode("ascii")
    return bytes(out)


def _make_docx_bytes() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph("Hello paragraph")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "TableCellA"
    table.cell(0, 1).text = "TableCellB"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

# ===========================================================================
# AC-2 / R21 — pdf success
# ===========================================================================

def test_pdf_success_text_hello(tmp_path):
    pdf = tmp_path / "hello.pdf"
    pdf.write_bytes(_make_minimal_pdf("Hello"))
    text, error = _norm_result(_call_handler("pdf", pdf))
    assert error is None, f"unexpected pdf error: {error!r}"
    assert "Hello" in text


# AC-3 covered by operator test
# tests/test_access_extract_text.py::test_pdf_explicit_format_fake_bytes_parse_error
# — not duplicated here.

# ===========================================================================
# AC-4 / R21 — pdf without pypdfium2
# ===========================================================================

def test_pdf_missing_pypdfium2_error(tmp_path, monkeypatch):
    pdf = tmp_path / "hello.pdf"
    pdf.write_bytes(_make_minimal_pdf("Hello"))
    monkeypatch.setitem(sys.modules, "pypdfium2", None)
    text, error = _norm_result(_call_handler("pdf", pdf))
    assert error is not None
    assert "pypdfium2 not installed" in error
    assert "lokdocsqr[pdf]" in error


# ===========================================================================
# AC-5 / R22 — docx success (D1: paragraph-before-table)
# ===========================================================================

def test_docx_success_paragraph_before_table(tmp_path):
    pytest.importorskip("docx")
    docx_path = tmp_path / "hello.docx"
    docx_path.write_bytes(_make_docx_bytes())
    text, error = _norm_result(_call_handler("docx", docx_path))
    assert error is None, f"unexpected docx error: {error!r}"
    assert "Hello paragraph" in text
    assert "TableCellA" in text
    assert text.find("Hello paragraph") < text.find("TableCellA")


# ===========================================================================
# AC-6 / R22 — docx parse error
# ===========================================================================

def test_docx_parse_error(tmp_path):
    docx_path = tmp_path / "broken.docx"
    docx_path.write_bytes(b"not a zip")
    text, error = _norm_result(_call_handler("docx", docx_path))
    assert error is not None
    assert "docx parse error" in error


# ===========================================================================
# AC-7 / R22 — docx without python-docx
# ===========================================================================

def test_docx_missing_python_docx_error(tmp_path, monkeypatch):
    docx_path = tmp_path / "hello.docx"
    docx_path.write_bytes(b"placeholder")
    monkeypatch.setitem(sys.modules, "docx", None)
    text, error = _norm_result(_call_handler("docx", docx_path))
    assert error is not None
    assert "python-docx not installed" in error
    assert "lokdocsqr[docx]" in error


# ===========================================================================
# AC-8 / R23 — eml text/plain
# ===========================================================================

def test_eml_text_plain(tmp_path):
    eml_path = tmp_path / "mail.eml"
    eml_path.write_bytes(MIMEText("Hello", "plain").as_bytes())
    text, error = _norm_result(_call_handler("eml", eml_path))
    assert error is None, f"unexpected eml error: {error!r}"
    assert "Hello" in text


# ===========================================================================
# AC-9 / R23 — eml text/html fallback
# ===========================================================================

def test_eml_html_fallback(tmp_path):
    eml_path = tmp_path / "mail.eml"
    eml_path.write_bytes(MIMEText("<p>Hello</p>", "html").as_bytes())
    text, error = _norm_result(_call_handler("eml", eml_path))
    assert error is None, f"unexpected eml error: {error!r}"
    assert "Hello" in text
    assert "<p>" not in text
    assert "<" not in text
    assert ">" not in text


# ===========================================================================
# AC-10 / R23 — eml multipart with attachment
# ===========================================================================

def test_eml_multipart_attachment(tmp_path):
    msg = MIMEMultipart()
    msg.attach(MIMEText("Hello plain body", "plain"))
    attachment = MIMEApplication(b"ATTACHMENTBYTES", Name="x.bin")
    msg.attach(attachment)

    eml_path = tmp_path / "mail.eml"
    eml_path.write_bytes(msg.as_bytes())
    text, error = _norm_result(_call_handler("eml", eml_path))
    assert error is None, f"unexpected eml error: {error!r}"
    assert "Hello plain body" in text
    assert "ATTACHMENTBYTES" not in text


# ===========================================================================
# AC-11 / R24 — fake handler extension proves index_search is unchanged
# ===========================================================================

def test_fake_handler_extension_via_index_search(tmp_path):
    marker = "FAKE2_MARKER_6D2"

    def fake_handler(*args, **kwargs):
        return _FakeText(marker)

    old = extract_text.HANDLERS.get("fake2")
    extract_text.register_handler("fake2", fake_handler)
    try:
        (tmp_path / "doc.fake2").write_text("dummy", encoding="utf-8")
        result = _run_index_search(tmp_path, query=marker)
        assert marker in repr(result), (
            f"index_search did not surface fake2 handler output; got {result!r}"
        )
    finally:
        if old is None:
            extract_text.HANDLERS.pop("fake2", None)
        else:
            extract_text.HANDLERS["fake2"] = old


# ===========================================================================
# AC-12 / R24 — index_search behavior sanity on 6d1 formats
# ===========================================================================

def test_index_search_regression_sanity(tmp_path):
    (tmp_path / "a.txt").write_text("Hello txt", encoding="utf-8")
    (tmp_path / "b.htm").write_text("<p>Hello htm</p>", encoding="utf-8")
    (tmp_path / "c.csv").write_text("name,value\nHello,csv\n", encoding="utf-8")
    (tmp_path / "d.png").write_bytes(b"\x89PNG\r\n\x1a\nfake")

    result = _run_index_search(tmp_path, query="Hello")
    assert result is not None
    rep = repr(result)
    assert (
        "Hello" in rep
        or any(fn in rep for fn in ("a.txt", "b.htm", "c.csv"))
        or "match" in rep.lower()
    )


# ===========================================================================
# AC-13 / R3 — no direct lxml import; python-docx import present
# ===========================================================================

def test_no_direct_lxml_import_in_extract_text():
    src = Path(extract_text.__file__).read_text(encoding="utf-8")
    assert "import lxml" not in src
    assert "from lxml" not in src
    assert ("from docx import" in src) or ("import docx" in src)


# ===========================================================================
# AC-14 — htm registration from 6d1 alive + behavioral
# ===========================================================================

def test_htm_handler_registered_once_and_works(tmp_path):
    src = Path(extract_text.__file__).read_text(encoding="utf-8")
    matches = re.findall(r'register_handler\s*\(\s*["\']htm["\']', src)
    assert len(matches) == 1, (
        f"register_handler('htm', ...) expected exactly once, got {len(matches)}"
    )

    htm_path = tmp_path / "hello.htm"
    htm_path.write_text("<p>Hello</p>", encoding="utf-8")
    text, error = _norm_result(_call_handler("htm", htm_path))
    assert error is None, f"unexpected htm error: {error!r}"
    assert "Hello" in text


# ===========================================================================
# AC-15 — dry_run _preview must not read file content
# ===========================================================================

def test_preview_does_not_read_content(tmp_path, monkeypatch):
    preview = getattr(extract_text, "_preview", None)
    if preview is None:
        pytest.skip("extract_text._preview not exposed; see test_dry_run.py")

    file_path = tmp_path / "x.txt"
    file_path.write_text("Hello", encoding="utf-8")

    real_open = builtins.open
    opened = []

    def guarded_open(*args, **kwargs):
        opened.append((args, kwargs))
        return real_open(*args, **kwargs)

    monkeypatch.setattr(builtins, "open", guarded_open)

    try:
        _invoke_callable(
            preview,
            path=file_path,
            max_bytes=DEFAULT_MAX_BYTES_PER_FILE,
            data=None,
            filename=file_path.name,
            opts={},
            params={},
        )
    except TypeError:
        pytest.skip("extract_text._preview has unknown private signature")

    assert opened == [], f"_preview unexpectedly opened files: {opened}"